"""Office-format -> markdown converters, feeding the same splitter pipeline as PDFs."""
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation


def _heading_level(style_name: str):
    """'Heading 2' -> 2, 'Title' -> 1, anything else -> None."""
    if not style_name:
        return None
    name = style_name.strip().lower()
    if name == "title":
        return 1
    if name.startswith("heading"):
        tail = name.replace("heading", "").strip()
        if tail.isdigit():
            return min(int(tail), 6)
    return None


def _table_to_markdown(table) -> str:
    rows = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ").replace("|", "\\|") for cell in row.cells]
        rows.append("| " + " | ".join(cells) + " |")

    if not rows:
        return ""

    width = len(table.rows[0].cells)
    separator = "| " + " | ".join(["---"] * width) + " |"
    return "\n".join([rows[0], separator] + rows[1:])


def _iter_docx_blocks(document):
    """Yield paragraphs and tables in document order (python-docx exposes them separately)."""
    body = document.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def docx_to_markdown(file_path: str) -> str:
    document = DocxDocument(file_path)
    parts = []

    for block in _iter_docx_blocks(document):
        if isinstance(block, Table):
            md = _table_to_markdown(block)
            if md:
                parts.append(md)
            continue

        text = block.text.strip()
        if not text:
            continue

        level = _heading_level(block.style.name if block.style else None)
        if level:
            parts.append(f"{'#' * level} {text}")
        elif (block.style.name or "").lower().startswith("list"):
            parts.append(f"- {text}")
        else:
            parts.append(text)

    return "\n\n".join(parts)


def _slide_title(slide) -> str:
    if slide.shapes.title is not None and slide.shapes.title.text.strip():
        return slide.shapes.title.text.strip()
    return ""


def pptx_to_markdown(file_path: str) -> str:
    presentation = Presentation(file_path)
    parts = []

    for index, slide in enumerate(presentation.slides, start=1):
        title = _slide_title(slide)
        parts.append(f"## Slide {index}" + (f": {title}" if title else ""))

        title_shape = slide.shapes.title
        for shape in slide.shapes:
            if shape == title_shape:
                continue

            if shape.has_table:
                md = _table_to_markdown(shape.table)
                if md:
                    parts.append(md)
                continue

            if not shape.has_text_frame:
                continue

            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if not text:
                    continue
                indent = "  " * max(paragraph.level, 0)
                parts.append(f"{indent}- {text}")

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"**Speaker notes:** {notes}")

    return "\n\n".join(parts)
